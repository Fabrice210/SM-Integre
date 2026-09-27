"""Écriture des tableaux : Excel (openpyxl) et CSV compatible Excel (comme le front)."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .tabular import Table

HEAD_FILL = PatternFill("solid", fgColor="E6F2EB")  # th du printDoc du front
BRAND = "17693F"


def to_xlsx(table: Table, org_name: str, stamp: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:31] or "Export"
    ws.append([table.title.replace("_", " ")])
    ws["A1"].font = Font(bold=True, size=14, color=BRAND)
    ws.append([f"{org_name} — généré le {stamp}"])
    ws["A2"].font = Font(italic=True, color="6D7772")
    ws.append([])
    ws.append(table.headers)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for row in table.rows:
        ws.append(row)
    for i, h in enumerate(table.headers, start=1):
        values = [str(h)] + [str(r[i - 1]) for r in table.rows]
        width = max(len(line) for v in values for line in (v.split("\n") or [""]))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A5"
    if table.headers:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(table.headers))}{max(4, ws.max_row)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(table: Table) -> bytes:
    """Même format que exportTable(…, 'xls') du front : BOM, « ; », CRLF, cellules entre guillemets."""
    buf = io.StringIO()
    buf.write("﻿" + ";".join(table.headers) + "\r\n")
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
    for row in table.rows:
        w.writerow(["" if v is None else str(v) for v in row])
    return buf.getvalue().encode("utf-8")
